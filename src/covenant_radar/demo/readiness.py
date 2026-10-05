"""Fail-fast presentation checks and reproducible, explicitly illustrative intake assets."""

from __future__ import annotations

# Console output is the public interface of this preflight command.
# ruff: noqa: T201
import argparse
import json
from datetime import UTC, date, datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path

SANCTION_CLAUSE = (
    "Debt service coverage ratio (DSCR) shall not fall below 1.25 times, "
    "tested quarterly, effective from 2026-04-01. A cure period of 30 days applies."
)
ASSET_NAMES = ("sanction-letter.pdf", "sanction-letter.docx", "readiness.json")


def check_live_gemini() -> None:
    """Exercise the actual extraction and grounded memo paths using synthetic facts."""
    from covenant_radar.ai import create_provider
    from covenant_radar.ai.client import InMemoryModelCallWriter, ModelClient
    from covenant_radar.ai.intake import propose_candidates
    from covenant_radar.ai.memo import draft_memo
    from covenant_radar.ai.shapes import CatalogueAction
    from covenant_radar.config.settings import get_settings
    from covenant_radar.domain.intake.candidates import CandidateLine, ClauseCandidate
    from covenant_radar.domain.memo import MemoRecord, MemoRecords, RecordReference
    from covenant_radar.services.memo import MemoAssemblyService

    settings = get_settings()
    if settings.ai.provider != "gemini":
        raise RuntimeError("Live presentation checks require Gemini and GEMINI_API_KEY.")
    provider = create_provider(settings.ai)
    try:
        client = ModelClient(
            provider, model=settings.ai.model, model_calls=InMemoryModelCallWriter()
        )
        line = CandidateLine(
            page_number=1, start_offset=0, end_offset=len(SANCTION_CLAUSE), text=SANCTION_CLAUSE
        )
        candidate = ClauseCandidate(
            start_page=1,
            start_offset=0,
            end_page=1,
            end_offset=len(SANCTION_CLAUSE),
            text=SANCTION_CLAUSE,
            matched_rules=("ratio:dscr",),
            lines=(line,),
        )
        proposal = propose_candidates((candidate,), client)[0]
        if not (
            proposal.parseable
            and proposal.definition_ref == "dscr"
            and proposal.threshold == Decimal("1.25")
            and proposal.direction == "min"
            and proposal.frequency == "quarterly"
        ):
            raise RuntimeError("Gemini extraction did not reproduce the known sample terms.")
        records = MemoRecords(
            situation=MemoRecord(
                RecordReference("triage", "demo-check"),
                {"situation": "Projected pressure requires review."},
            ),
            covenant_position=MemoRecord(
                RecordReference("forecast", "demo-check"),
                {
                    "ratio_name": "Debt service coverage",
                    "value": Decimal("1.25"),
                    "threshold": Decimal("1.10"),
                    "headroom": Decimal("0.15"),
                    "probability": Decimal("0.42"),
                    "confidence": Decimal("0.88"),
                    "crossing_date": date(2026, 10, 15),
                },
            ),
            drivers=(
                MemoRecord(RecordReference("driver", "demo-check"), {"name": "Cash-flow pressure"}),
            ),
            evidence=(
                MemoRecord(
                    RecordReference("evidence", "demo-check"), {"citation": "EV-DEMO", "count": 3}
                ),
            ),
            recommendations=(
                MemoRecord(
                    RecordReference("intervention", "demo-check"),
                    {
                        "code": "CREDIT-REDUCE",
                        "role_tag": "credit",
                        "text": "Review and reduce funded exposure.",
                    },
                ),
            ),
        )
        slots = MemoAssemblyService().assemble(records)
        draft_memo(
            slots,
            (
                CatalogueAction(
                    id="CREDIT-REDUCE", role_tag="credit", text="Review and reduce funded exposure."
                ),
            ),
            client,
        )
    finally:
        close = getattr(provider, "close", None)
        if callable(close):
            close()


def prepare_assets(output: Path) -> dict[str, object]:
    """Check real PDF/DOCX dependencies, scanner and native text extraction."""
    from docx import Document
    from pypdf import PdfReader
    from weasyprint import HTML  # type: ignore[import-untyped]

    from covenant_radar.security.uploads import UploadGuard

    output.mkdir(parents=True, exist_ok=True)
    pdf = HTML(
        string=f"""<!doctype html><html lang="en"><meta charset="utf-8">
    <style>@page {{ size: A4; margin: 24mm; }} body {{ font: 12pt sans-serif; line-height: 1.6; }}
    h1 {{ font-size: 24pt; color: #42283d; }} small {{ color: #555; }}</style>
    <h1>Illustrative sanction letter</h1>
    <p><strong>Covenant Radar presentation sample</strong></p>
    <p>This document is synthetic.
    It is not a sanction or an allegation about any listed company.</p>
    <h2>Financial covenant</h2><p>{SANCTION_CLAUSE}</p>
    <h2>Review control</h2><p>A credit officer must verify the source clause and proposed terms.
    Covenant registration requires the applicable approval workflow.</p>
    <small>For local demonstration only. Borrower and facility are selected at upload.</small>
    </html>"""
    ).write_pdf()
    text = "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(pdf)).pages)
    if "1.25" not in text or "quarterly" not in text:
        raise RuntimeError("The sample PDF does not contain readable covenant terms.")
    guard = UploadGuard()
    guard.validate("sanction-letter.pdf", "application/pdf", pdf)
    (output / ASSET_NAMES[0]).write_bytes(pdf)
    document = Document()
    document.add_heading("Illustrative sanction letter", 0)
    document.add_paragraph("Synthetic presentation sample; not an actual sanction.")
    document.add_paragraph(SANCTION_CLAUSE)
    stream = BytesIO()
    document.save(stream)
    docx = stream.getvalue()
    guard.validate(
        ASSET_NAMES[1],
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        docx,
    )
    (output / ASSET_NAMES[1]).write_bytes(docx)
    return {
        "checked_at": datetime.now(UTC).isoformat(),
        "pdf": "passed",
        "docx": "passed",
        "upload_validation": "passed",
        "native_pdf_text": "passed",
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--live", action="store_true")
    args = parser.parse_args()
    try:
        result = prepare_assets(args.output)
        if args.live:
            check_live_gemini()
        result["gemini"] = "extraction and grounded memo passed" if args.live else "offline replay"
        (args.output / ASSET_NAMES[2]).write_text(json.dumps(result, indent=2), encoding="utf-8")
    except Exception as error:
        # Provider exception chains may contain credentials. Emit only the safe class name.
        print(
            f"Presentation preflight failed ({type(error).__name__}). "
            "Check Gemini access/model, native PDF libraries and upload scanner before presenting."
        )
        return 2
    print(
        "Presentation preflight passed: PDF, DOCX, upload validation and "
        + ("live Gemini extraction/memo." if args.live else "explicit offline mode.")
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
